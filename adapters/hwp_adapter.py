"""한글(.hwp/.hwpx) 파일을 '서식 그대로' 가명처리하는 어댑터.

- 원본은 절대 열지 않는다. 먼저 AI용 경로로 복사한 뒤, 그 복사본만 연다.
- 치환 대상 찾기는 core.masker.mask_text 가 한다.
- 변경 목록의 '원래 값 → 바뀐 값'을 긴 것부터 한글의 '모두 바꾸기'로 적용한다.
  (본문·표·머리말·꼬리말·메모 전부에 적용된다.)
- 문서 정보(제목/지은이/주제/키워드)는 모두 비운다.
- 저장 후 다시 열어서 글자를 꺼내 core.checker.check_text 로 남은 개인정보를 검사한다.
- 한글은 백그라운드로 실행하고, 끝나면 반드시 닫는다. (Windows + 한글 설치 환경 전용)
"""
import re
import shutil
import sys
from pathlib import Path
from xml.sax.saxutils import unescape as _xml_unescape

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from core.checker import check_text
from core.masker import mask_text

HWP_EXTS = {".hwp", ".hwpx"}
HWP_WINDOWS_ONLY_MSG = "한글 파일은 한글이 설치된 Windows 에서만 처리할 수 있어요"
# 동명이인이 있을 때 보여줄 경고 (원본 형식 유지 방식은 전역 바꾸기라 부정확할 수 있음)
DUPLICATE_WARNING = "⚠ 같은 이름이 있어 ‘원래 형식’은 정확하지 않을 수 있어요. ‘AI에 올릴 글’을 이용해 주세요."


class HwpWindowsOnlyError(RuntimeError):
    """Windows 가 아닌 환경에서 한글(.hwp/.hwpx) 파일을 처리하려 할 때 발생."""


# HWPML 에서 머리말/꼬리말/메모 글자를 뽑기 위한 정규식
_HWPML_HEADER = re.compile(r"<HEADER\b[^>]*>(.*?)</HEADER>", re.S)
_HWPML_FOOTER = re.compile(r"<FOOTER\b[^>]*>(.*?)</FOOTER>", re.S)
_HWPML_MEMO = re.compile(r'<FIELDBEGIN\b[^>]*\bType="Memo"[^>]*>(.*?)</FIELDBEGIN>', re.S)
_HWPML_CHAR = re.compile(r"<CHAR>(.*?)</CHAR>", re.S)


def _blocks(ml, pattern):
    """HWPML 에서 pattern 으로 잡은 블록 안의 <CHAR> 글자들을 목록으로."""
    out = []
    for block in pattern.findall(ml):
        text = _xml_unescape("".join(_HWPML_CHAR.findall(block))).strip()
        if text:
            out.append(text)
    return out


def _collect_text(hwp):
    """열린 문서의 모든 글자를 하나의 문자열로 모은다.

    본문·표는 get_text_file(UNICODE), 머리말·꼬리말·메모는 HWPML 에서 따로 뽑아 붙인다.
    (본문 문단은 한 줄로 유지되어 '성명+학번'이 같은 줄에 오므로 동명이인 판별에 유리하다.)
    """
    main = hwp.get_text_file("UNICODE", "") or ""
    ml = hwp.GetTextFile("HWPML2X", "") or ""
    extras = _blocks(ml, _HWPML_HEADER) + _blocks(ml, _HWPML_FOOTER) + _blocks(ml, _HWPML_MEMO)
    if extras:
        return main + "\r\n" + "\r\n".join(extras)
    return main


def _dedupe_changes(raw_changes):
    """mask_text 의 변경목록을 '원래 값' 기준으로 합치고 긴 원래값부터 정렬한다.

    Args:
        raw_changes: [[원래, 바뀐, 종류, 위치], ...]

    Returns:
        [(원래, 바뀐, 종류, 적용횟수), ...]  (원래==바뀐 은 제외, 긴 원래값 우선)
    """
    order, seen = [], {}
    for original, new, kind, _pos in raw_changes:
        if original == new:
            continue
        if original in seen:
            seen[original]["count"] += 1
            continue
        seen[original] = {"new": new, "kind": kind, "count": 1}
        order.append(original)
    items = [(o, seen[o]["new"], seen[o]["kind"], seen[o]["count"]) for o in order]
    # 긴 원래값을 먼저 바꿔야 짧은 값(예: 성 뺀 이름)이 긴 값(성명 전체)을 깨뜨리지 않는다
    items.sort(key=lambda t: (-len(t[0]), t[0]))
    return items


def mask_hwp(src_path, dst_path, dictionary, confirm_duplicate=None):
    """원본을 dst 로 복사한 뒤 복사본만 열어 서식 그대로 가명처리하고 저장한다.

    Args:
        src_path: 원본 한글 파일 (열지 않고 복사만 함)
        dst_path: 가명처리 결과를 저장할 경로 (원본과 같은 확장자 권장)
        dictionary: core.masker.load_dictionary 로 만든 학생사전
        confirm_duplicate: 동명이인이 있을 때 호출되는 콜백(인자 없음).
            True 반환 시 계속 진행, False 반환 시 중단. None 이면 그냥 진행.

    Returns:
        dict(status, changes, findings, duplicates)
          - status: "done"(완료) / "aborted"(동명이인으로 사용자가 중단)
          - changes: [(원래, 바뀐, 종류, 적용횟수), ...]
          - findings: 저장 후 재검사에서 남은 개인정보(Finding) 목록 (aborted 면 빈 목록)
          - duplicates: 동명이인 포함 여부
    """
    if sys.platform != "win32":
        raise HwpWindowsOnlyError(HWP_WINDOWS_ONLY_MSG)

    src_path = Path(src_path)
    dst_path = Path(dst_path)
    dst_path.parent.mkdir(parents=True, exist_ok=True)

    # 1) 원본을 복사한다. 원본은 한글로 열지 않는다. (복사본만 연다)
    shutil.copy2(src_path, dst_path)

    from pyhwpx import Hwp

    hwp = None
    try:
        hwp = Hwp(new=True, visible=False)   # 백그라운드 실행
        hwp.set_visible(False)
        hwp.open(str(dst_path))              # 복사본만 연다

        # 2) 문서 전체 글자(표·머리말·꼬리말·메모 포함)를 꺼내 mask_text 로 변경 목록을 얻는다
        text = _collect_text(hwp)
        _, raw_changes = mask_text(text, dictionary)
        items = _dedupe_changes(raw_changes)
        duplicates = any(kind == "동명이인확인" for _, _, kind, _ in items)

        # [주의] 동명이인이 있으면 바꾸기 전에 사용자에게 계속할지 물어본다
        if duplicates and confirm_duplicate is not None and not confirm_duplicate():
            return {"status": "aborted", "changes": items, "findings": [], "duplicates": duplicates}

        # 3) 변경 목록을 긴 원래값부터 한글 '모두 바꾸기'로 적용 (본문·표·머리말·꼬리말·메모 모두)
        for original, new, _kind, _count in items:
            hwp.find_replace_all(original, new)

        # 4) 문서 정보(제목/지은이/주제/키워드)를 모두 비운다
        summary = hwp.XHwpDocuments.Item(0).XHwpSummaryInfo
        for attr in ("Title", "Author", "Subject", "Keywords"):
            try:
                setattr(summary, attr, "")
            except Exception:
                pass

        # 저장 (복사본 경로에 같은 형식으로 그대로 저장됨)
        hwp.save(save_if_dirty=False)

        # 5) 저장 후 다시 열어서 글자를 꺼내 검사한다
        hwp.open(str(dst_path))
        findings = check_text(_collect_text(hwp), dictionary)

        return {"status": "done", "changes": items, "findings": findings, "duplicates": duplicates}
    finally:
        # 성공/실패와 무관하게 한글은 반드시 닫는다
        if hwp is not None:
            try:
                hwp.quit()
            except Exception:
                pass


def apply_extra_names_hwp(path, words):
    """가명처리된 .hwp 에서 사용자가 고른 단어를 [이름N] 으로 바꾼다 (한글 설치 필요)."""
    if sys.platform != "win32":
        raise HwpWindowsOnlyError(HWP_WINDOWS_ONLY_MSG)
    uniq = list(dict.fromkeys(w for w in words if w))
    name_map = {f"[이름{i}]": w for i, w in enumerate(uniq, start=1)}
    if not name_map:
        return [], {}
    from pyhwpx import Hwp
    hwp = None
    items = []
    try:
        hwp = Hwp(new=True, visible=False)
        hwp.set_visible(False)
        hwp.open(str(path))
        for code, word in name_map.items():
            hwp.find_replace_all(word, code)
            items.append((word, code, "이름의심확정", 1))
        hwp.save(save_if_dirty=False)
    finally:
        if hwp is not None:
            try:
                hwp.quit()
            except Exception:
                pass
    return items, name_map


def write_hwp_change_log(changes, path, name_map=None):
    """보관용 변경기록 엑셀: 원래 값 / 바뀐 값 / 종류 / 적용 횟수.

    changes: [(원래, 바뀐, 종류, 적용횟수), ...]
    """
    from adapters.excel_adapter import append_name_map_sheet
    wb = Workbook()
    ws = wb.active
    ws.title = "변경기록"
    ws.append(["원래 값", "바뀐 값", "종류", "적용 횟수"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for original, new, kind, count in changes:
        ws.append([original, new, kind, count])
    for row in ws.iter_rows():
        for cell in row:   # 앞 0·'=' 로 시작하는 값이 변하지 않게 모두 글자로 저장
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.number_format = "@"
    for i in range(1, 5):
        letter = get_column_letter(i)
        longest = max((len(str(c.value or "")) for c in ws[letter]), default=8)
        ws.column_dimensions[letter].width = min(max(longest + 3, 10), 60)
    ws.freeze_panes = "A2"
    append_name_map_sheet(wb, name_map)
    wb.save(path)
