"""여러 형식의 파일을 md(글)로 바꾼 뒤 가명처리하는 어댑터.

- .pdf .docx .xlsx .pptx .csv : markitdown 으로 md 변환 (LLM·이미지 설명·외부 서비스·플러그인 안 씀)
                               워드 머리말·꼬리말·메모·각주, 피피티 슬라이드 노트는 따로 뽑아 '부가 정보'로 덧붙임
- .hwpx                       : ZIP+XML 을 읽어 md 로 바꾼다 (한글 프로그램 불필요)
- .hwp                        : pyhwpx 로 한글을 백그라운드로 열어 임시 HTML 로 저장 후 markitdown 으로 변환
                               (머리말·꼬리말·메모는 따로 뽑아 '부가 정보'로 덧붙임, Windows+한글 설치 환경 전용)
- .txt .md                    : 그대로 읽음
- 변환한 원본 md(실명이 있는 글)는 메모리에서만 다루고 파일로 저장하지 않는다.
- 글자 치환은 core.masker.mask_text 가 한다.
"""
import re
from pathlib import Path
from xml.sax.saxutils import unescape as _xml_unescape

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from core.masker import mask_text

CONVERT_EXTS = {".pdf", ".docx", ".xlsx", ".pptx", ".csv"}
TEXT_EXTS = {".txt", ".md"}
HWP_EXTS = {".hwp", ".hwpx"}
SUPPORTED_EXTS = CONVERT_EXTS | TEXT_EXTS | HWP_EXTS

# 맥 등 Windows 가 아닌 환경에서 한글 파일을 넣었을 때 보여줄 안내 문구
HWP_WINDOWS_ONLY_MSG = "한글 파일은 한글이 설치된 Windows 에서만 처리할 수 있어요"


class HwpWindowsOnlyError(RuntimeError):
    """Windows 가 아닌 환경에서 한글(.hwp/.hwpx) 파일을 처리하려 할 때 발생."""

SCAN_MIN_CHARS = 10   # 공백 뺀 글자가 이보다 적은 쪽은 스캔된 쪽으로 본다

# 엑셀을 표로 바꾼 글에서: 표의 칸 하나가 숫자 하나뿐인 경우
_CELL_FLOAT = re.compile(r"(?<=\|)(\s*)(\d{5,})\.0(\s*)(?=\|)")     # 10305.0 → 10305
_CELL_LOST_ZERO = re.compile(r"(?<=\|)(\s*)(10\d{8})(\s*)(?=\|)")   # 1012345678 (앞 0이 빠진 번호)


def decode_text(raw):
    for enc in ("utf-8-sig", "cp949"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ValueError("글자를 읽지 못했어요. UTF-8 또는 한글(CP949) 파일만 돼요.")


def find_scanned_pages(path):
    """PDF에서 추출되는 글자가 거의 없는 쪽 번호(1부터) 목록. PDF는 읽기만 한다."""
    import pymupdf   # requirements.txt 의 pymupdf
    with pymupdf.open(path) as doc:
        return [
            i for i, page in enumerate(doc, start=1)
            if len("".join(page.get_text().split())) < SCAN_MIN_CHARS
        ]


def to_markdown(path):
    """파일을 md 글로 바꿔서 돌려준다 (메모리에서만 다루고 파일로 저장하지 않음)."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext in TEXT_EXTS:
        return decode_text(path.read_bytes())
    if ext == ".hwpx":
        from adapters.hwpx_adapter import hwpx_to_markdown
        return hwpx_to_markdown(path)
    if ext in HWP_EXTS:
        return _hwp_to_markdown(path)
    if ext not in CONVERT_EXTS:
        raise ValueError(f"이 형식은 쓸 수 없어요: {ext}")

    from markitdown import MarkItDown
    # llm_client 등을 주지 않고 플러그인도 끄면 이미지 설명(LLM)·외부 서비스를 쓰지 않는다.
    converter = MarkItDown(enable_plugins=False)
    # convert() 는 http:// 주소도 받아들이므로, 로컬 파일만 받는 convert_local() 을 쓴다.
    text = converter.convert_local(str(path), keep_data_uris=False).text_content
    extra = office_extra_markdown(path)
    if extra:
        return text.rstrip() + "\n\n" + extra
    return text


def office_extra_markdown(path):
    """워드 머리말·꼬리말·메모·각주, 피피티 슬라이드 노트를 md 조각으로 붙인다. 없으면 빈 글."""
    ext = Path(path).suffix.lower()
    if ext == ".docx":
        return _docx_extra_info(path)
    if ext == ".pptx":
        return _pptx_extra_info(path)
    return ""


def _format_extra_sections(sections):
    lines = []
    for title, items in sections:
        if not items:
            continue
        lines.append(f"### {title}")
        lines.extend(f"- {item}" for item in items)
        lines.append("")
    if not lines:
        return ""
    return "## 부가 정보\n\n" + "\n".join(lines).rstrip() + "\n"


def _docx_part_texts(part, tag):
    """docx XML 파트에서 w:t 글자를 모아 목록으로 돌린다."""
    from docx.oxml.ns import qn
    texts = []
    for node in part.element.findall(qn(tag)):
        joined = "".join((t.text or "") for t in node.iter(qn("w:t"))).strip()
        if joined:
            texts.append(joined)
    return texts


def _docx_hf_texts(header_or_footer):
    texts = [p.text.strip() for p in header_or_footer.paragraphs if p.text.strip()]
    for table in header_or_footer.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t:
                    texts.append(t)
    return texts


def _docx_extra_info(path):
    try:
        from docx import Document
    except ImportError:
        return ""
    try:
        doc = Document(str(path))
    except Exception:
        return ""
    headers, footers = [], []
    for section in doc.sections:
        headers.extend(_docx_hf_texts(section.header))
        footers.extend(_docx_hf_texts(section.footer))
    comments, footnotes, endnotes = [], [], []
    for rel in doc.part.rels.values():
        kind = (rel.reltype or "").lower()
        try:
            part = rel.target_part
        except Exception:
            continue
        if "comments" in kind:
            comments.extend(_docx_part_texts(part, "w:comment"))
        elif "footnotes" in kind:
            footnotes.extend(_docx_part_texts(part, "w:footnote"))
        elif "endnotes" in kind:
            endnotes.extend(_docx_part_texts(part, "w:endnote"))
    return _format_extra_sections([
        ("머리말", list(dict.fromkeys(headers))),
        ("꼬리말", list(dict.fromkeys(footers))),
        ("메모", list(dict.fromkeys(comments))),
        ("각주", list(dict.fromkeys(footnotes))),
        ("미주", list(dict.fromkeys(endnotes))),
    ])


def _pptx_extra_info(path):
    try:
        from pptx import Presentation
    except ImportError:
        return ""
    try:
        prs = Presentation(str(path))
    except Exception:
        return ""
    notes, comments = [], []
    for i, slide in enumerate(prs.slides, 1):
        if getattr(slide, "has_notes_slide", False):
            try:
                text = (slide.notes_slide.notes_text_frame.text or "").strip()
            except Exception:
                text = ""
            if text:
                notes.append(f"{i}번째 슬라이드: {text}")
        try:
            for c in getattr(slide, "comments", []) or []:
                t = (getattr(c, "text", "") or "").strip()
                if t:
                    comments.append(f"{i}번째 슬라이드: {t}")
        except Exception:
            pass
    return _format_extra_sections([
        ("슬라이드 노트", notes),
        ("메모", comments),
    ])


# ── 한글(.hwp/.hwpx) 전용 처리 ──────────────────────────────────────────────
# HWPML 에서 머리말/꼬리말/메모 글자만 뽑아내기 위한 정규식
_HWPML_HEADER = re.compile(r"<HEADER\b[^>]*>(.*?)</HEADER>", re.S)
_HWPML_FOOTER = re.compile(r"<FOOTER\b[^>]*>(.*?)</FOOTER>", re.S)
_HWPML_MEMO = re.compile(r'<FIELDBEGIN\b[^>]*\bType="Memo"[^>]*>(.*?)</FIELDBEGIN>', re.S)
_HWPML_CHAR = re.compile(r"<CHAR>(.*?)</CHAR>", re.S)


def _hwpml_block_texts(ml, pattern):
    """HWPML 에서 pattern 으로 잡은 블록들 안의 <CHAR> 글자를 꺼내 목록으로 돌려준다."""
    texts = []
    for block in pattern.findall(ml):
        joined = "".join(_HWPML_CHAR.findall(block))
        text = _xml_unescape(joined).strip()
        if text:
            texts.append(text)
    return texts


def _hwp_extra_info(ml):
    """머리말·꼬리말·메모 글자를 모아 '부가 정보' md 조각으로 만든다. 없으면 빈 문자열."""
    sections = [
        ("머리말", _hwpml_block_texts(ml, _HWPML_HEADER)),
        ("꼬리말", _hwpml_block_texts(ml, _HWPML_FOOTER)),
        ("메모", _hwpml_block_texts(ml, _HWPML_MEMO)),
    ]
    lines = []
    for title, items in sections:
        if not items:
            continue
        lines.append(f"### {title}")
        lines.extend(f"- {item}" for item in items)
        lines.append("")
    if not lines:
        return ""
    return "## 부가 정보\n\n" + "\n".join(lines).rstrip() + "\n"


def _hwp_to_markdown(path):
    """한글 파일을 md 로 바꾼다 (Windows+한글 설치 환경 전용).

    1) pyhwpx 로 한글을 백그라운드로 열어 문서를 임시 HTML 로 저장한다. (원본은 읽기만 함)
    2) 그 HTML 을 markitdown 으로 md 로 바꾼다. (표는 md 표로)
    3) 머리말·꼬리말·메모 글자는 HWPML 에서 따로 뽑아 '부가 정보'로 맨 아래에 붙인다.

    임시 HTML 에는 실명이 들어 있으므로 try/finally 로 처리가 끝나면(오류가 나도) 반드시 지운다.
    한글 프로그램도 끝나면 반드시 닫는다.
    """
    import sys
    import shutil
    import tempfile

    if sys.platform != "win32":
        raise HwpWindowsOnlyError(HWP_WINDOWS_ONLY_MSG)

    from markitdown import MarkItDown
    from pyhwpx import Hwp

    path = Path(path)
    tmp_dir = Path(tempfile.mkdtemp(prefix="hwp_md_"))
    html_path = tmp_dir / "본문.html"
    hwp = None
    try:
        # 1) 한글 백그라운드 실행 → 원본 열기(읽기 전용으로만 사용) → 임시 HTML 저장
        hwp = Hwp(new=True, visible=False)
        hwp.set_visible(False)
        hwp.open(str(path))
        hwp.save_as(str(html_path), format="HTML")
        # 3) 머리말/꼬리말/메모 추출용 HWPML 을 메모리로 받아둔다
        ml = hwp.GetTextFile("HWPML2X", "")
        # 변환 전에 한글을 먼저 닫아 파일 핸들을 푼다
        hwp.quit()
        hwp = None

        # 2) 임시 HTML → md (표 포함). LLM·외부 서비스·플러그인 안 씀.
        converter = MarkItDown(enable_plugins=False)
        body_md = converter.convert_local(str(html_path), keep_data_uris=False).text_content
        extra_md = _hwp_extra_info(ml)
    finally:
        # 오류가 나도: 한글은 반드시 닫고, 실명이 든 임시 HTML(및 임시폴더)은 반드시 지운다.
        if hwp is not None:
            try:
                hwp.quit()
            except Exception:
                pass
        shutil.rmtree(tmp_dir, ignore_errors=True)

    if extra_md:
        return body_md.rstrip() + "\n\n" + extra_md
    return body_md


def prepare_excel_markdown(text):
    """엑셀에서 바꾼 md 정리: 숫자 칸의 '.0'을 없애고, 앞 0이 빠진 10자리 번호에 0을 붙인다.

    (정리한 글, 0을 붙인 자리의 글자 번호 집합)을 돌려준다.
    """
    text = _CELL_FLOAT.sub(lambda m: f"{m.group(1)}{m.group(2)}{m.group(3)}", text)

    pieces, added, last, offset = [], set(), 0, 0
    for m in _CELL_LOST_ZERO.finditer(text):
        start = m.start(2)
        pieces.append(text[last:start])
        pieces.append("0")
        added.add(start + offset)
        offset += 1
        last = start
    pieces.append(text[last:])
    return "".join(pieces), added


def mask_markdown(text, dictionary, added_zero=frozenset()):
    """md 글을 가명처리한다. (가명처리된 글, 변경 목록 [[줄, 원래 값, 바뀐 값, 종류], ...])."""
    masked, ch = mask_text(text, dictionary)
    changes = []
    for original, new, kind, pos in ch:
        if pos in added_zero:   # 우리가 붙인 0은 원래 값에서 뺀다
            original = original[1:]
        changes.append([text.count("\n", 0, pos) + 1, original, new, kind])
    return masked, changes


def extra_name_map(words, start=1):
    """확정한 이름 단어들을 [이름N] 코드로 매긴다. ({코드: 단어}, 순서 유지)"""
    uniq = list(dict.fromkeys(w for w in words if w))
    return {f"[이름{i}]": w for i, w in enumerate(uniq, start=start)}


def apply_extra_names(text, words, start=1):
    """이름 단어를 [이름N] 으로 바꾼다.

    (바뀐 글, 변경목록 [[줄, 원래, 바뀐, 종류], ...], {코드: 단어}) 를 돌려준다.
    """
    name_map = extra_name_map(words, start=start)
    if not name_map:
        return text, [], {}
    word_to_code = {w: c for c, w in name_map.items()}
    # 꼬리말/조사(씨, 에게, 와 ...)가 바로 뒤에 붙어도 이름을 바꾸도록 오른쪽 경계는 두지 않는다
    alt = "|".join(re.escape(w) for w in sorted(word_to_code, key=lambda w: (-len(w), w)))
    pat = re.compile(rf"(?<![가-힣])(?:{alt})")
    changes = []

    def repl(m):
        w = m.group(0)
        changes.append([text.count("\n", 0, m.start()) + 1, w, word_to_code[w], "이름의심확정"])
        return word_to_code[w]

    return pat.sub(repl, text), changes, name_map


def write_md_change_log(changes, path, name_map=None):
    """보관용 변경기록 엑셀: 줄, 원래 값, 바뀐 값, 종류."""
    from adapters.excel_adapter import append_name_map_sheet
    wb = Workbook()
    ws = wb.active
    ws.title = "변경기록"
    ws.append(["줄", "원래 값", "바뀐 값", "종류"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for line, original, new, kind in changes:
        ws.append([line, original, new, kind])
    for row in ws.iter_rows(min_row=2):
        for cell in row[1:]:
            if isinstance(cell.value, str):
                cell.data_type = "s"     # '=' 로 시작하는 글이 수식이 되지 않게
            cell.number_format = "@"     # 앞 0이 사라지지 않게
    for i in range(1, 5):
        letter = get_column_letter(i)
        longest = max((len(str(c.value or "")) for c in ws[letter]), default=8)
        ws.column_dimensions[letter].width = min(max(longest + 3, 10), 60)
    ws.freeze_panes = "A2"
    append_name_map_sheet(wb, name_map)
    wb.save(path)
