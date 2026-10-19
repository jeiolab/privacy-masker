"""samples/가짜_명렬표.xlsx 의 학생 정보로 테스트용 가짜 파일 2개를 만든다.

- samples/가짜_성적.xlsx      : 학번(숫자), 합계(SUM 수식), 학생연락처(숫자라 앞 0이 빠짐)
- samples/가짜_상담기록.xlsx  : 시트 이름에 실명, 긴 문장 칸, 셀 메모(코멘트)

make_sample.py 를 먼저 실행해 명렬표가 있어야 한다.   실행: python3 make_sample2.py
"""

import random
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parent
ROSTER = ROOT / "samples" / "가짜_명렬표.xlsx"
GRADES = ROOT / "samples" / "가짜_성적.xlsx"
COUNSEL = ROOT / "samples" / "가짜_상담기록.xlsx"
RANDOM_SEED = 42


# ---------------------------------------------------------------- 명렬표 읽기
def load_roster():
    ws = load_workbook(ROSTER, read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    return [dict(zip(header, r)) for r in rows[1:] if r[header.index("성명")]]


def given_name(name):
    return name[1:]  # 이 명렬표에는 한 글자 성만 있다


def digits(phone):
    return phone.replace("-", "")


def phone_fmt(phone, style):
    """같은 번호를 서로 다른 모양으로 쓴다."""
    d = digits(phone)
    a, b, c = d[:3], d[3:7], d[7:]
    return {"hyphen": f"{a}-{b}-{c}", "plain": d, "space": f"{a} {b} {c}", "dot": f"{a}.{b}.{c}"}[style]


# ---------------------------------------------------------------- 가짜_성적.xlsx
def make_grades(students):
    rng = random.Random(RANDOM_SEED)
    wb = Workbook()
    ws = wb.active
    ws.title = "성적"
    ws.append(["학번", "성명", "국어", "수학", "영어", "합계", "학생연락처"])
    for c in ws[1]:
        c.font = Font(bold=True)

    for i, s in enumerate(students, start=2):
        ws.append([
            int(s["학번"]),                       # 숫자 형식 (10305)
            s["성명"],
            rng.randint(50, 100),
            rng.randint(50, 100),
            rng.randint(50, 100),
            f"=SUM(C{i}:E{i})",
            int(digits(s["학생연락처"])),          # 숫자라서 맨 앞 0이 빠짐 (1012345678)
        ])
        ws.cell(i, 1).number_format = "0"
        ws.cell(i, 7).number_format = "0"

    for col, w in zip("ABCDEFG", (10, 10, 8, 8, 8, 8, 14)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(GRADES)


# ---------------------------------------------------------------- 가짜_상담기록.xlsx
def parent_title(s):
    return "아버님" if s["관계"] == "부" else "어머님"


def counsel_lines(s):
    """학생 한 명에 대한 상담 5줄. 성명, 성 뺀 이름+조사, 보호자명, 전화번호, 학번이 섞인다."""
    n, g, gd = s["성명"], given_name(s["성명"]), s["보호자명"]
    return [
        f"{n} 학생이 수업 시간에 자주 졸아 면담함. {g}이는 요즘 늦게 잔다고 했고, "
        f"학번 {s['학번']} 기준으로 이번 달 지각 3회 확인.",
        f"{gd} {parent_title(s)}과 유선 상담({phone_fmt(s['보호자연락처'], 'hyphen')}). "
        f"{g}이가 집에서 게임을 오래 한다고 걱정하심.",
        f"{g}아, 다음 주까지 과제를 내기로 약속함. 학생 연락처 {phone_fmt(s['학생연락처'], 'plain')} 로 문자 발송.",
        f"{gd} 보호자가 {phone_fmt(s['보호자연락처'], 'dot')} 로 다시 전화 줌. "
        f"{n}은 방과후 수업 신청을 희망함.",
        f"[{s['학번']}] {n} / {g} 학생 담임 확인: 2학기 상담 재실시 예정. "
        f"연락처 {phone_fmt(s['학생연락처'], 'space')}.",
    ]


def memo_lines(class_students):
    """학급 메모 5줄 (여러 학생이 섞임). 동명이인·일반 단어 이름도 들어간다."""
    a, b, c, d, e = (class_students[i] for i in (0, 1, 3, 5, 7))
    dup = next(s for s in class_students if s["성명"] == "정하윤")
    nara = next(s for s in class_students if s["성명"] == "이나라")
    return [
        f"{a['성명']}({a['학번']}) 결석 사유 확인 필요. 보호자 {a['보호자명']} 연락처 {phone_fmt(a['보호자연락처'], 'hyphen')}.",
        f"{b['성명']}과 {given_name(d['성명'])}이가 쉬는 시간에 다툼. {b['학번']}, {d['학번']} 모두 지도함.",
        f"{dup['성명']}({dup['학번']}) 학생 보호자 {dup['보호자명']} 면담 희망. 학생 번호 {phone_fmt(dup['학생연락처'], 'plain')}.",
        f"{nara['성명']} 학생은 발표를 잘함. 나라이 발전하려면 인재가 필요하다고 말해 줌. 보호자 {nara['보호자명']}.",
        f"{e['성명']}는 체험학습 신청서 제출. {given_name(e['성명'])}이를 칭찬해 주기. "
        f"연락처 {phone_fmt(e['학생연락처'], 'dot')}, 주소 {e['주소']}.",
    ]


def write_lines(ws, lines):
    ws.append(["일시", "내용"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, text in enumerate(lines, start=1):
        ws.append([f"2026-03-{i + 1:02d}", text])
        ws.cell(i + 1, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(i + 1, 1).alignment = Alignment(vertical="top")
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 80
    ws.freeze_panes = "A2"


def make_counsel(students):
    target = next(s for s in students if s["성명"] == "김준서")  # 실명이 시트 이름에 들어간다
    class1 = students[:10]

    wb = Workbook()
    ws1 = wb.active
    ws1.title = f"{target['성명']} 상담"          # 예: "김준서 상담"
    write_lines(ws1, counsel_lines(target))

    # 셀 메모(코멘트)에도 학생 이름을 넣는다
    g = given_name(target["성명"])
    note = Comment(
        f"{target['성명']} 학생 상담 중 {g}이가 울먹임. 보호자 {target['보호자명']} 님께 "
        f"{phone_fmt(target['보호자연락처'], 'hyphen')} 로 연락 필요.",
        "담임",
    )
    note.width, note.height = 320, 120
    ws1["B3"].comment = note

    ws2 = wb.create_sheet("1학년 1반 메모")
    write_lines(ws2, memo_lines(class1))
    wb.save(COUNSEL)


def main():
    if not ROSTER.exists():
        raise SystemExit("samples/가짜_명렬표.xlsx 가 없습니다. 먼저 python3 make_sample.py 를 실행하세요.")
    students = load_roster()
    make_grades(students)
    make_counsel(students)
    print(f"만들었습니다: {GRADES.relative_to(ROOT)}, {COUNSEL.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
