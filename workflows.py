"""화면(app.py)이 부르는 작업 모음. 화면 코드는 하나도 없다.

core/, adapters/ 의 기존 함수를 그대로 불러서, mask_md.py / mask_excel.py / build_dict.py /
restore.py 와 같은 순서로 일을 하되, 터미널에 출력하는 대신 결과를 돌려준다.
(터미널용 실행 파일들은 그대로 두었다.)
"""
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

import build_dict
from adapters.excel_adapter import (
    Change, apply_extra_names_workbook, load_name_map, mask_workbook, restore_workbook,
    safe_filename, write_change_log,
)
from adapters.hwp_adapter import write_hwp_change_log
from adapters.md_adapter import (
    CONVERT_EXTS, TEXT_EXTS, apply_extra_names, decode_text, find_scanned_pages, mask_markdown,
    prepare_excel_markdown, to_markdown, write_md_change_log,
)
from core.names import leftover_person_names
from core import config
from core.checker import add_ignored_word as _add_ignored_word
from core.checker import check_file, check_md_file, check_text, find_name_suspects
from core.checker import load_ignored_words as _load_ignored_words
from core.checker import remove_ignored_words as _remove_ignored_words
from core.masker import load_dictionary, mask_text
from core.restorer import unmask_text

BASE_DIR = config.user_data_dir()   # 문서\개인정보가리개_데이터 (경로는 config 한 곳)
DICT_PATH = config.dict_path()
IGNORE_PATH = config.ignore_words_path()
AI_DIR = config.ai_dir()
KEEP_DIR = config.keep_dir()
RESTORE_DIR = config.restore_dir()
OUTPUT_DIRS = {"AI용": AI_DIR, "보관용": KEEP_DIR, "복원": RESTORE_DIR}  # config.CLEANUP_DAYS 와 짝이 되는 폴더들
SPECIAL_KINDS = ("사전에없음", "동명이인확인")
HWP_EXTS = {".hwp", ".hwpx"}
HWP_MESSAGE = "한글(.hwp) 파일은 한글이 설치된 Windows 에서만 처리할 수 있어요"
MASK_EXTS = CONVERT_EXTS | TEXT_EXTS | HWP_EXTS | {".xls"}
PREVIEW_LIMIT = 200_000   # 엑셀 결과를 글로 보여 줄 때 최대 글자 수


# ---------------------------------------------------------------- 사전
def dictionary_info():
    """(학생 수, 교직원 수, 수정 날짜 'YYYY-MM-DD'). 사전이 없으면 None. 읽을 수 없으면 오류를 그대로 낸다."""
    if not DICT_PATH.exists():
        return None
    d = load_dictionary(DICT_PATH)
    students = len(d.students)
    staff = len(getattr(d, "staff", []))
    modified = datetime.fromtimestamp(DICT_PATH.stat().st_mtime).strftime("%Y-%m-%d")
    return students, staff, modified


def get_dictionary():
    if not DICT_PATH.exists():
        raise FileNotFoundError("학생사전 없음 - 먼저 사전을 만들어 주세요")
    return load_dictionary(DICT_PATH)


def add_ignored_word(word, dictionary=None, path=None):
    """단어를 data/무시단어.txt 에 저장한다. 사전 실명이면 거부 문구를 돌려준다."""
    return _add_ignored_word(word, dictionary or get_dictionary(), path=path or IGNORE_PATH)


def load_ignored_words(path=None):
    return _load_ignored_words(path or IGNORE_PATH)


def remove_ignored_words(words, path=None):
    return _remove_ignored_words(words, path or IGNORE_PATH)


# ---------------------------------------------------------------- output 폴더 정리
# 중요: 아래 함수들은 output/ 아래 결과 폴더(AI용·보관용·복원)의 '파일'만 지운다.
#       data/ (사전·백업) 는 절대 건드리지 않는다.
def _delete_files(folder, older_than_days=None):
    """folder 안의 파일을 지운다. older_than_days 를 주면 그 일수보다 오래된 파일만 지운다. (지운 개수)"""
    if not folder.exists() or not folder.is_dir():
        return 0
    cutoff = None if older_than_days is None else datetime.now().timestamp() - older_than_days * 86400
    removed = 0
    for item in folder.iterdir():
        if not item.is_file():
            continue  # 혹시 모를 하위 폴더는 손대지 않는다
        try:
            if cutoff is not None and item.stat().st_mtime >= cutoff:
                continue  # 아직 안 지난 파일은 남겨 둔다
            item.unlink()
            removed += 1
        except OSError:
            pass  # 쓰는 중이거나 잠긴 파일은 조용히 넘어간다
    return removed


def cleanup_old_outputs():
    """프로그램 시작 때 호출: config.CLEANUP_DAYS 일수가 지난 결과 파일을 지운다. (지운 개수)"""
    total = 0
    for name, folder in OUTPUT_DIRS.items():
        days = config.CLEANUP_DAYS.get(name)
        if days is None:
            continue
        total += _delete_files(folder, older_than_days=days)
    return total


def clear_all_outputs():
    """output/ 의 결과 파일을 모두 비운다. (data/ 는 건드리지 않음) (지운 개수)"""
    total = 0
    for folder in OUTPUT_DIRS.values():
        total += _delete_files(folder, older_than_days=None)
    return total


@dataclass
class BuildResult:
    count: int
    found: list
    missing: list
    warnings: list
    path: Path
    mode: str = "새로 만들기"          # "새로 만들기" / "사전 업데이트"
    stats: dict = None                  # 업데이트일 때 {"kept","updated","added","missing"}
    backup: Path = None                 # 업데이트일 때 백업 파일 경로


def build_dictionary(src, kind="학생"):
    """명렬표/교직원 명단 엑셀로 data/학생사전.xlsx 를 만들거나 업데이트한다(추가 모드).

    kind="학생"  : 학년 전체 명렬표(다른 반 포함)를 학생사전 시트에 추가 모드로 반영
    kind="교직원": 교직원 명단을 '교직원' 시트에 추가 모드로 반영
    둘 다 한 번 붙은 코드 번호는 절대 바꾸지 않는다.
    """
    if kind == "교직원":
        return build_staff_dictionary(src)
    src = Path(src)
    if src.suffix.lower() == ".xls":
        with tempfile.TemporaryDirectory() as tmp:
            return build_dictionary(convert_xls_to_xlsx(src, tmp))
    if src.suffix.lower() != ".xlsx":
        raise ValueError("학생 명렬표는 엑셀 파일만 넣을 수 있어요.")
    out_path = config.dict_path()
    if src.resolve() == out_path.resolve():
        raise ValueError("지금 쓰는 명단 파일은 고를 수 없어요. 원본 엑셀을 골라 주세요.")

    warnings = []
    new_students, mapping = build_dict.collect_students(src, warnings)
    if mapping is None:
        raise ValueError("엑셀 위쪽에서 이름 칸을 찾지 못했어요.")
    found = [f for f in config.COLUMN_CANDIDATES if f in mapping]
    missing = [f for f in config.COLUMN_CANDIDATES if f not in mapping]
    if "성명" not in mapping:
        raise ValueError("이름 칸을 찾지 못해 명단을 넣지 못했어요.")
    warnings = [f"{f} 열을 찾지 못해 건너뜀" for f in missing] + warnings
    if not new_students:
        raise ValueError("학생이 한 명도 없어요. 명렬표를 확인해 주세요.")

    existing = build_dict.load_existing_dict(out_path)
    if existing is None:
        # 사전이 없으면 지금처럼 새로 만든다
        build_dict.assign_codes(new_students, mapping, warnings)
        build_dict.add_cautions(new_students, warnings)
        build_dict.write_dict(new_students, mapping, out_path)
        return BuildResult(len(new_students), found, missing, warnings, out_path)

    # 사전이 있으면 바꾸기 전에 백업하고, 코드 번호를 유지한 채 업데이트한다
    backup = build_dict.backup_dict(out_path)
    students, stats = build_dict.merge_students(existing, new_students, mapping, warnings)
    build_dict.add_cautions(students, warnings)
    build_dict.apply_missing_notes(students, warnings)
    students.sort(key=build_dict._num_key)
    build_dict.write_dict_incremental(students, out_path)
    summary = (f"그대로 둔 사람 {stats['kept']}명 / 새로 넣은 사람 {stats['added']}명 / "
               f"바뀐 정보 {stats['updated']}명 / 명렬표에 없는 사람 {stats['missing']}명")
    warnings.insert(0, f"명단을 이어서 넣었어요. {summary}")
    return BuildResult(len(students), found, missing, warnings, out_path,
                       mode="사전 업데이트", stats=stats, backup=backup)


def build_staff_dictionary(src):
    """교직원 명단 엑셀(.xlsx / .xls)로 data/학생사전.xlsx 의 '교직원' 시트를 추가 모드로 반영한다.

    기존 교사 코드 번호(교사001 ...)는 절대 바뀌지 않는다. 학생사전 시트는 그대로 둔다.
    """
    src = Path(src)
    if src.suffix.lower() == ".xls":
        with tempfile.TemporaryDirectory() as tmp:
            return build_staff_dictionary(convert_xls_to_xlsx(src, tmp))
    if src.suffix.lower() != ".xlsx":
        raise ValueError("교직원 명단은 엑셀 파일만 넣을 수 있어요.")
    out_path = config.dict_path()
    if src.resolve() == out_path.resolve():
        raise ValueError("지금 쓰는 명단 파일은 고를 수 없어요. 원본 엑셀을 골라 주세요.")

    _, names = build_dict.load_staff_names(src)
    if not names:
        raise ValueError("교직원 이름을 한 명도 찾지 못했어요. 이름 칸을 확인해 주세요.")

    warnings = []
    existing = build_dict.load_existing_staff(out_path)   # 파일 없으면 None, 시트 없으면 []
    backup = None
    if existing is None:
        merged, stats = build_dict.merge_staff([], names, warnings)
        mode = "새로 만들기"
    else:
        backup = build_dict.backup_dict(out_path)
        merged, stats = build_dict.merge_staff(existing, names, warnings)
        mode = "업데이트"
    build_dict.write_staff_sheet(merged, out_path)

    summary = (f"그대로 둔 사람 {stats['kept']}명 / 새로 넣은 사람 {stats['added']}명 / "
               f"명단에 없는 사람 {stats['missing']}명")
    how = "새로" if mode == "새로 만들기" else "이어서"
    warnings.insert(0, f"교직원 명단을 {how} 넣었어요. {summary}")
    return BuildResult(len(merged), ["성명"], [], warnings, out_path,
                       mode=f"교직원 {mode}", stats=stats, backup=backup)


# ---------------------------------------------------------------- .xls → .xlsx (임시)
def convert_xls_to_xlsx(src, tmp_dir):
    """.xls 를 임시 폴더 안의 같은 이름 .xlsx 로 바꾼다. (값만 옮김. 임시 폴더는 부르는 쪽이 지운다)"""
    try:
        import xlrd
    except ImportError:
        raise RuntimeError("예전 엑셀(.xls)을 읽지 못했어요.")
    src = Path(src)
    book = xlrd.open_workbook(str(src))
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in book.sheets():
        ws = wb.create_sheet(title=sheet.name[:31] or "시트")
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype == xlrd.XL_CELL_NUMBER:
                    row.append(int(cell.value) if float(cell.value).is_integer() else cell.value)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    row.append(bool(cell.value))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                    row.append(None)
                else:
                    row.append(cell.value)
            ws.append(row)
    if not wb.worksheets:
        wb.create_sheet("시트")
    dst = Path(tmp_dir) / f"{src.stem}.xlsx"
    wb.save(dst)
    return dst


# ---------------------------------------------------------------- 가리기
@dataclass
class MaskResult:
    ok: bool                       # 결과 파일을 만들었는가
    headline: str                  # 검사 결과 한 줄
    clean: bool                    # 남은 개인정보 없음
    text: str = ""                 # 결과 내용 (미리보기·복사용)
    ai_path: Path = None
    ai_paths: list = field(default_factory=list)   # 여러 파일이면 연 결과 파일 전부
    keep_path: Path = None
    total_changes: int = 0
    review: list = field(default_factory=list)      # (종류, 위치, 내용) 확인 필요 목록
    findings: list = field(default_factory=list)    # (위치, 종류, 남은 글자)
    suspects: list = field(default_factory=list)    # 이름 의심 (Suspect: location, word, context)
    name_map: dict = field(default_factory=dict)    # '이 단어도 가리기'로 바꾼 {코드: 단어}

    @property
    def can_copy(self):
        return self.ok and self.clean


def mask_file(path, mode, dictionary, extra_names=()):
    """mode: 'md'(AI용 md) 또는 'xlsx'(엑셀·한글 원본 형식 유지).

    extra_names: 사용자가 '이 단어도 가리기'로 고른 이름 단어 목록. [이름N] 으로 추가 치환한다.
    """
    src = Path(path)
    ext = src.suffix.lower()
    if ext not in MASK_EXTS:
        raise ValueError(f"이 형식은 쓸 수 없어요: {ext}")
    if ext == ".xls":
        with tempfile.TemporaryDirectory() as tmp:     # 변환본은 임시로만 쓰고 바로 지움
            converted = convert_xls_to_xlsx(src, tmp)
            if mode == "xlsx":
                return _mask_to_excel(converted, dictionary, extra_names)
            return _mask_to_md(converted, dictionary, extra_names)
    if mode == "xlsx":
        if ext == ".xlsx":
            return _mask_to_excel(src, dictionary, extra_names)
        if ext in HWP_EXTS:
            return _mask_to_hwp(src, dictionary, extra_names)
        raise ValueError("‘원래 형식’은 엑셀(.xlsx)과 한글(.hwp, .hwpx)만 쓸 수 있어요.")
    return _mask_to_md(src, dictionary, extra_names)


def mask_files(paths, mode, dictionary, extra_names=()):
    """여러 파일을 순서대로 가명처리하고 결과를 하나로 합친다."""
    results, errors = [], []
    for path in paths:
        try:
            results.append(mask_file(path, mode, dictionary, extra_names))
        except Exception as e:
            errors.append((Path(path), str(e)))
    if not results:
        msg = "; ".join(f"{p.name}: {err}" for p, err in errors) or "가릴 파일이 없어요."
        raise ValueError(msg)
    return _combine_mask_results(results, errors)


def _combine_mask_results(results, errors=()):
    if len(results) == 1 and not errors:
        return results[0]
    ok = all(r.ok for r in results) and not errors
    clean = all(r.clean for r in results) and not errors
    texts, paths = [], []
    review, findings, suspects = [], [], []
    seen, name_map = set(), {}
    for r in results:
        paths.extend(r.ai_paths or ([r.ai_path] if r.ai_path else []))
        name = r.ai_path.name if r.ai_path else "결과"
        texts.append(f"## {name}\n\n{r.text}".rstrip())
        review.extend(r.review)
        findings.extend(r.findings)
        name_map.update(r.name_map)
        for s in r.suspects:
            if s.word in seen:
                continue
            seen.add(s.word)
            suspects.append(s)
    for p, err in errors:
        review.append(("오류", p.name, err))
    n, fail_n = len(results), len(errors)
    if fail_n:
        headline = f"⚠ {fail_n}개를 가리지 못했어요 · {n}개는 가렸어요"
    elif not clean:
        headline = f"⚠ {n}개 중 개인정보가 남았어요. AI에 올리지 마세요"
    else:
        headline = f"✅ {n}개 파일, 남은 개인정보가 없어요"
    return MaskResult(
        ok=ok, clean=clean, headline=headline, text="\n\n".join(texts),
        ai_path=paths[0] if paths else None, ai_paths=paths, keep_path=results[0].keep_path,
        total_changes=sum(r.total_changes for r in results),
        review=review, findings=findings, suspects=suspects, name_map=name_map,
    )


def _staff_notice(dictionary):
    if dictionary is not None and not getattr(dictionary, "staff", None):
        return [("안내", "", "교직원 명단이 없어요. 선생님 이름이 남을 수 있어요")]
    return []


def _masked_stem(src, dictionary):
    stem, name_changes = mask_text(src.stem, dictionary)
    return safe_filename(stem), name_changes


def _mask_to_md(src, dictionary, extra_names=()):
    ext = src.suffix.lower()
    if ext == ".pdf":
        scanned = find_scanned_pages(src)
        if scanned:
            pages = ", ".join(map(str, scanned))
            return MaskResult(ok=False, clean=False, headline=f"⚠ 그림으로만 된 {pages}쪽은 가릴 수 없어요")

    text = to_markdown(src)            # 실명이 든 원본 md: 메모리에서만 다룸 (파일로 저장 안 함)
    if not text.strip():
        raise ValueError("파일에서 글을 읽지 못했어요.")
    added_zero = frozenset()
    if ext in (".xlsx", ".csv"):
        text, added_zero = prepare_excel_markdown(text)
    masked, changes = mask_markdown(text, dictionary, added_zero)

    # 사전에 없어도 칸·글에 남은 사람 이름 + 사용자가 고른 단어를 [이름N] 으로 바꾼다
    name_map = {}
    extra_words = leftover_person_names(masked, dictionary) + list(extra_names)
    extra_words = list(dict.fromkeys(w for w in extra_words if w))
    if extra_words:
        masked, extra_changes, name_map = apply_extra_names(masked, extra_words)
        changes += extra_changes

    stem, name_changes = _masked_stem(src, dictionary)
    changes = [["(파일 이름)", o, n, k] for o, n, k, _ in name_changes] + changes
    ai_path, keep_path = AI_DIR / f"{stem}_가명.md", KEEP_DIR / f"{stem}_변경기록.xlsx"
    if src.resolve() in (ai_path.resolve(), keep_path.resolve()):
        raise ValueError("고른 파일이 결과 파일과 같아요. 다른 파일을 골라 주세요.")
    AI_DIR.mkdir(parents=True, exist_ok=True)
    KEEP_DIR.mkdir(parents=True, exist_ok=True)
    ai_path.write_text(masked, encoding="utf-8")
    write_md_change_log(changes, keep_path, name_map)

    findings = [(f.location, f.kind, f.value) for f in check_md_file(ai_path, dictionary)]
    review = [
        (kind, line if isinstance(line, str) else f"{line}번째 줄", f"{o} → {n}")
        for line, o, n, kind in changes if kind in SPECIAL_KINDS
    ]
    suspects = find_name_suspects(masked, dictionary)
    return _finish(ai_path, keep_path, masked, len(changes), review, findings, suspects, name_map,
                   dictionary)


def _mask_to_excel(src, dictionary, extra_names=()):
    stem, name_changes = _masked_stem(src, dictionary)
    ai_path, keep_path = AI_DIR / f"{stem}_가명.xlsx", KEEP_DIR / f"{stem}_변경기록.xlsx"
    if src.resolve() in (ai_path.resolve(), keep_path.resolve()):
        raise ValueError("고른 파일이 결과 파일과 같아요. 다른 파일을 골라 주세요.")
    AI_DIR.mkdir(parents=True, exist_ok=True)
    KEEP_DIR.mkdir(parents=True, exist_ok=True)

    changes, warnings = mask_workbook(src, ai_path, dictionary)
    changes = [Change("(파일 이름)", "", "-", o, n, k) for o, n, k, _ in name_changes] + changes

    text = workbook_text(ai_path)
    extra_words = leftover_person_names(text, dictionary) + list(extra_names)
    extra_words = list(dict.fromkeys(w for w in extra_words if w))
    name_map = {}
    if extra_words:
        extra_changes, name_map = apply_extra_names_workbook(ai_path, extra_words)
        changes += extra_changes
        text = workbook_text(ai_path)
    write_change_log(changes, keep_path, name_map)
    findings = [(f"{f.sheet} / {f.location}", f.kind, f.value) for f in check_file(ai_path, dictionary)]
    review = [(c.kind, f"{c.sheet} / {c.cell}", f"{c.original} → {c.new}") for c in changes if c.kind in SPECIAL_KINDS]
    review += [("경고", "", w) for w in warnings]
    suspects = find_name_suspects(text, dictionary)
    return _finish(ai_path, keep_path, text, len(changes), review, findings, suspects, name_map,
                   dictionary)


def _mask_to_hwp(src, dictionary, extra_names=()):
    """한글 원본 형식을 유지한 채 가명처리한다. .hwpx 는 한컴 없이, .hwp 는 한글 설치 필요."""
    from adapters.hwp_adapter import apply_extra_names_hwp, mask_hwp
    from adapters.hwpx_adapter import apply_extra_names_hwpx, hwpx_text, hwpx_to_markdown, mask_hwpx

    ext = src.suffix.lower()
    stem, name_changes = _masked_stem(src, dictionary)
    ai_path, keep_path = AI_DIR / f"{stem}_가명{ext}", KEEP_DIR / f"{stem}_변경기록.xlsx"
    if src.resolve() in (ai_path.resolve(), keep_path.resolve()):
        raise ValueError("고른 파일이 결과 파일과 같아요. 다른 파일을 골라 주세요.")
    AI_DIR.mkdir(parents=True, exist_ok=True)
    KEEP_DIR.mkdir(parents=True, exist_ok=True)

    if ext == ".hwpx":
        result = mask_hwpx(src, ai_path, dictionary)
    else:
        result = mask_hwp(src, ai_path, dictionary)
    if result["status"] == "aborted":
        raise ValueError("동명이인이 있어 원래 형식으로는 가릴 수 없어요. ‘AI에 올릴 글’을 이용해 주세요.")

    items = [(o, n, k, 1) for o, n, k, _ in name_changes] + list(result["changes"])

    if ext == ".hwpx":
        preview = hwpx_text(ai_path)
    else:
        try:
            preview = to_markdown(ai_path)
        except Exception:
            preview = ""
    extra_words = leftover_person_names(preview, dictionary) + list(extra_names)
    extra_words = list(dict.fromkeys(w for w in extra_words if w))
    name_map = {}
    if extra_words:
        extra_fn = apply_extra_names_hwpx if ext == ".hwpx" else apply_extra_names_hwp
        extra_items, name_map = extra_fn(ai_path, extra_words)
        items += extra_items

    write_hwp_change_log(items, keep_path, name_map)

    if ext == ".hwpx":
        text = hwpx_to_markdown(ai_path)
        findings = [(f.location, f.kind, f.value) for f in check_text(hwpx_text(ai_path), dictionary)]
    else:
        findings = [(f.location, f.kind, f.value) for f in result.get("findings", [])]
        try:
            text = to_markdown(ai_path)
        except Exception:
            text = "\n".join(f"{o} → {n}" for o, n, _k, _c in items)

    review = [
        (kind, "문서", f"{o} → {n}")
        for o, n, kind, _c in items if kind in SPECIAL_KINDS
    ]
    if result.get("duplicates"):
        review.append(("동명이인", "문서", "같은 이름이 있어 ‘원래 형식’은 정확하지 않을 수 있어요"))
    suspects = find_name_suspects(text, dictionary)
    return _finish(ai_path, keep_path, text, len(items), review, findings, suspects, name_map,
                   dictionary)


def _finish(ai_path, keep_path, text, total, review, findings, suspects=(), name_map=None,
            dictionary=None):
    review = list(review) + _staff_notice(dictionary)
    saved = f"저장: {ai_path.relative_to(BASE_DIR)}"
    if findings:
        headline = f"⚠ 개인정보가 {len(findings)}건 남았어요. AI에 올리지 마세요 ({saved})"
    else:
        headline = f"✅ 남은 개인정보가 없어요  ({saved})"
    return MaskResult(ok=True, clean=not findings, headline=headline, text=text, ai_path=ai_path,
                      ai_paths=[ai_path] if ai_path else [],
                      keep_path=keep_path, total_changes=total, review=review, findings=findings,
                      suspects=list(suspects), name_map=name_map or {})


def workbook_text(path):
    """가명처리된 엑셀을 미리보기용 글로 (시트마다 '## 이름' 아래에 한 줄씩)."""
    wb = load_workbook(path)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"## {ws.title}")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                parts.append(" | ".join("" if v is None else _cell_text(v) for v in row))
        parts.append("")
    text = "\n".join(parts)
    if len(text) > PREVIEW_LIMIT:
        text = text[:PREVIEW_LIMIT] + "\n… (이하 생략)"
    return text


def _cell_text(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return str(v)


# ---------------------------------------------------------------- 되돌리기
@dataclass
class RestoreFileResult:
    out_path: Path
    text: str | None      # .txt/.md 는 복원된 글, .xlsx 는 None
    changes: int
    warnings: list


def find_keep_log(src):
    """가명 파일 이름에서 짝이 되는 보관용 변경기록을 찾는다."""
    src = Path(src)
    stem = src.stem
    if stem.endswith("_가명"):
        stem = stem[:-3]
    from folder_mask import is_folder_output, keep_log_for   # 폴더 가리기(90_가림본)로 만든 파일은 원본 쪽 _가림기록/ 에 짝이 있다 (1.3.0)
    if is_folder_output(src):               # 폴더 가리기 결과면 그 짝만 쓴다 — 없으면 다른 작업의 같은 이름 기록을 쓰지 않고 None (강희 지적)
        return keep_log_for(src)
    for p in (KEEP_DIR / f"{stem}_변경기록.xlsx", src.parent / f"{stem}_변경기록.xlsx"):
        if p.exists():
            return p
    return None


def restore_file(path, dictionary, extra=None):
    """.txt / .md / .xlsx 를 복원해서 output/복원/ 에 저장 (restore.py 와 같은 방식).

    extra: '이 단어도 가리기'로 바꾼 {코드: 단어} 매핑 (txt/md 복원에서 [이름N] 을 되돌릴 때 쓴다).
    프로그램이 꺼진 뒤에는 보관용 변경기록의 이름대응 시트로 [이름N] 을 되돌린다.
    """
    src = Path(path)
    ext = src.suffix.lower()
    if ext not in (".txt", ".md", ".xlsx"):
        raise ValueError("되돌리기는 글 파일(.txt, .md)과 엑셀(.xlsx)만 돼요.")
    merged = {}
    keep = find_keep_log(src)
    if keep:
        merged.update(load_name_map(keep))
    merged.update(extra or {})
    RESTORE_DIR.mkdir(parents=True, exist_ok=True)
    stem = safe_filename(src.stem)
    if ext == ".xlsx":
        out = RESTORE_DIR / f"{stem}_복원.xlsx"
        if out.resolve() == src.resolve():
            raise ValueError("고른 파일이 결과 파일과 같아요. 다른 파일을 골라 주세요.")
        changes, warnings = restore_workbook(src, out, dictionary, extra=merged)
        return RestoreFileResult(out, None, len(changes), warnings)
    out = RESTORE_DIR / f"{stem}_복원.txt"
    restored, changes, warnings = unmask_text(decode_text(src.read_bytes()), dictionary, merged)
    out.write_text(restored, encoding="utf-8")
    return RestoreFileResult(out, restored, len(changes), warnings)
